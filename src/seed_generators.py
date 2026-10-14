import os
import zipfile
import json
import xml.etree.ElementTree as ET
try:
    from reportlab.pdfgen import canvas
except ImportError:
    canvas = None

try:
    from docx import Document
except ImportError:
    Document = None

try:
    from openpyxl import Workbook
except ImportError:
    Workbook = None

try:
    from pptx import Presentation
except ImportError:
    Presentation = None


# Runs used to write LLM-generated files straight into the TRACKED source tree.
# src/predefined_seeds/ held 36 files at b1b657e and had grown to 50 -- the extra
# 14 written by fuzzing runs (test_case_*.json). Every run shared the directory,
# so one run's seeds leaked into the next and the git tree went dirty during a
# campaign. TOFU_SEED_DIR gives each run its own; the default is unchanged so
# targets/fuzz_*.py and FuzzFileManager keep resolving "predefined_seed:<name>".
SEED_DIR = os.environ.get("TOFU_SEED_DIR", "src/predefined_seeds")

def create_pdf(content: str, filename: str) -> str:
    """
    Creates a valid PDF document with the specified content and saves it.
    Returns the reference format: 'predefined_seed:<filename>'
    """
    os.makedirs(SEED_DIR, exist_ok=True)
    if not filename.endswith(".pdf"):
        filename += ".pdf"
    filepath = os.path.join(SEED_DIR, filename)
    c = canvas.Canvas(filepath)
    y = 750
    for line in content.split("\n"):
        c.drawString(100, y, line)
        y -= 20
        if y < 50:
            c.showPage()
            y = 750
    c.save()
    return f"predefined_seed:{filename}"

def create_docx(content: str, filename: str) -> str:
    """
    Creates a valid Word (DOCX) document with the specified content and saves it.
    Returns the reference format: 'predefined_seed:<filename>'
    """
    os.makedirs(SEED_DIR, exist_ok=True)
    if not filename.endswith(".docx"):
        filename += ".docx"
    filepath = os.path.join(SEED_DIR, filename)
    doc = Document()
    for line in content.split("\n"):
        doc.add_paragraph(line)
    doc.save(filepath)
    return f"predefined_seed:{filename}"

def create_xlsx(sheet_data: list, filename: str) -> str:
    """
    Creates a valid Excel (XLSX) spreadsheet with the specified table data and saves it.
    sheet_data is a list of rows, where each row is a list of cell values.
    Returns the reference format: 'predefined_seed:<filename>'
    """
    os.makedirs(SEED_DIR, exist_ok=True)
    if not filename.endswith(".xlsx"):
        filename += ".xlsx"
    filepath = os.path.join(SEED_DIR, filename)
    wb = Workbook()
    ws = wb.active
    ws.title = "Fuzz Sheet"
    for r_idx, row in enumerate(sheet_data, start=1):
        if isinstance(row, dict):
            list_vals = [v for v in row.values() if isinstance(v, list)]
            if list_vals:
                row = list_vals[0]
            else:
                row = list(row.values())
        elif not isinstance(row, list):
            row = [row]
            
        for c_idx, val in enumerate(row, start=1):
            if isinstance(val, (dict, list)):
                val = json.dumps(val)
            ws.cell(row=r_idx, column=c_idx, value=val)
    wb.save(filepath)
    return f"predefined_seed:{filename}"

def create_pptx(content: str, filename: str) -> str:
    """
    Creates a valid PowerPoint (PPTX) presentation with the specified content and saves it.
    Returns the reference format: 'predefined_seed:<filename>'
    """
    os.makedirs(SEED_DIR, exist_ok=True)
    if not filename.endswith(".pptx"):
        filename += ".pptx"
    filepath = os.path.join(SEED_DIR, filename)
    prs = Presentation()
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title_shape = slide.shapes.title
    body_shape = slide.placeholders[1]
    title_shape.text = "Fuzz Seed"
    body_shape.text = content
    prs.save(filepath)
    return f"predefined_seed:{filename}"

def create_zip(files: dict, filename: str) -> str:
    """
    Creates a valid ZIP archive containing the specified file names and their text contents and saves it.
    files is a dictionary mapping sub-filename to its text content.
    Returns the reference format: 'predefined_seed:<filename>'
    """
    os.makedirs(SEED_DIR, exist_ok=True)
    if not filename.endswith(".zip"):
        filename += ".zip"
    filepath = os.path.join(SEED_DIR, filename)
    with zipfile.ZipFile(filepath, "w", zipfile.ZIP_DEFLATED) as z:
        for sub_name, sub_content in files.items():
            if not isinstance(sub_content, (str, bytes)):
                sub_content = json.dumps(sub_content)
            z.writestr(sub_name, sub_content)
    return f"predefined_seed:{filename}"

def create_xml(xml_content: str, filename: str) -> str:
    """
    Creates a valid XML file with the specified content and saves it.
    Returns the reference format: 'predefined_seed:<filename>'
    """
    os.makedirs(SEED_DIR, exist_ok=True)
    if not filename.endswith(".xml"):
        filename += ".xml"
    filepath = os.path.join(SEED_DIR, filename)
    root = ET.fromstring(xml_content)
    tree = ET.ElementTree(root)
    tree.write(filepath, encoding="utf-8", xml_declaration=True)
    return f"predefined_seed:{filename}"

def create_html(html_content: str, filename: str) -> str:
    """
    Creates a valid HTML file with the specified content and saves it.
    Returns the reference format: 'predefined_seed:<filename>'
    """
    os.makedirs(SEED_DIR, exist_ok=True)
    if not filename.endswith(".html"):
        filename += ".html"
    filepath = os.path.join(SEED_DIR, filename)
    with open(filepath, "w", encoding="utf-8") as f:
        f.write(html_content)
    return f"predefined_seed:{filename}"

def create_json(json_content: str, filename: str) -> str:
    """
    Creates a valid JSON file with the specified content and saves it.
    Returns the reference format: 'predefined_seed:<filename>'
    """
    os.makedirs(SEED_DIR, exist_ok=True)
    if not filename.endswith(".json"):
        filename += ".json"
    filepath = os.path.join(SEED_DIR, filename)
    data = json.loads(json_content)
    with open(filepath, "w", encoding="utf-8") as f:
        json.dump(data, f, indent=2)
    return f"predefined_seed:{filename}"
